"""Build the zonal deck (.pptx) — DustLedger, InnovateX Theme 2, Delhi zonal 10 Oct 2026.
Text mirrors ../zonal-pitch.md; figures come from make_figures.py (committed results only).
Run: cd deliverables/deck && python make_deck.py  → DustLedger_zonal.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
INK, ACCENT, GREY, WARN = RGBColor(0x1F, 0x29, 0x37), RGBColor(0x0F, 0x76, 0x6E), RGBColor(0x6B, 0x72, 0x80), RGBColor(0xB4, 0x53, 0x09)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, lines, size=20, color=INK, bold_first=False, space=6):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(ln, tuple):
            ln, sz, col, bold = ln
        else:
            sz, col, bold = size, color, (bold_first and i == 0)
        p.text = ln
        p.font.size, p.font.color.rgb, p.font.bold = Pt(sz), col, bold
        p.space_after = Pt(space)
    return tf


def slide(title, notes, kicker=None):
    s = prs.slides.add_slide(BLANK)
    if kicker:
        text(s, 0.6, 0.3, 12, 0.4, [(kicker, 14, ACCENT, True)])
    text(s, 0.6, 0.6, 12.2, 1.0, [(title, 32, INK, True)])
    s.notes_slide.notes_text_frame.text = notes
    return s


def source(s, line):
    text(s, 0.6, 6.95, 12.2, 0.4, [(line, 10, GREY, False)])


# 1 — Title
s = slide("DustLedger", "This winter Delhi will stop every construction site for six weeks, and restrict dusty work at all of them for three months, whether a site is clean or dirty. We're team ___, and we're building a way to tell them apart.", "InnovateX · Theme 2: Rising AQI in Modern Urban Cities · Delhi zonal, 10 Oct 2026")
text(s, 0.6, 1.7, 12, 2.4, [("Measure each construction site's own dust —", 30, INK, False), ("so Delhi can stop the dirty sites, not all of them.", 30, ACCENT, True), ("धूल का हिसाब  ·  Dhool ka hisaab", 20, GREY, False)])
text(s, 0.6, 5.2, 12, 1.4, [("Team: ______  ·  Team ID: ______  ·  Members & contacts: ______", 18, GREY, False)])

# 2 — The blanket ban
s = slide("This winter, Delhi will restrict every construction site", "Delhi's Winter Action Plan 2026-27 prohibits dust-generating demolition and outdoor civil works from November 1 to January 31, and bans all construction and demolition from December 10 to January 20. GRAP emergency stages add stops on top: in 2025, GRAP-3 halted construction for 16 days. These measures hit every site, whatever its dust control.", "The problem")
text(s, 0.6, 1.9, 4.0, 3.5, [("Nov 1 – Jan 31", 34, ACCENT, True), ("dust-generating demolition and outdoor civil works prohibited", 18, INK, False)])
text(s, 4.8, 1.9, 4.0, 3.5, [("Dec 10 – Jan 20", 34, ACCENT, True), ("total ban on construction and demolition", 18, INK, False)])
text(s, 9.0, 1.9, 3.9, 3.5, [("16 days", 34, ACCENT, True), ("GRAP-3 construction halt in 2025 (₹10,000 compensation announced)", 18, INK, False)])
text(s, 0.6, 4.6, 12, 1.5, [("Every site is treated the same — clean or dirty.", 24, INK, True)])
source(s, "Sources: Delhi Winter Action Plan 2026-27 (The News Mill, Sept 2026); Hindustan Times, Dec 2025. Full citations: research/landscape.md")

# 3 — Why
s = slide("A monitor can tell the air is bad — not which site made it bad", "Delhi measures air at city scale: 47 reference stations, going to 60, under a planning norm of one station per 25 to 50 square kilometres. A monitor can't say which of the sites around it is dusty. So enforcement falls back on checklists — is an anti-smog gun installed, is the self-audit uploaded, is there CCTV — and fines are flat per violation. And trust is broken: in October 2025, water tankers were filmed spraying a reference monitor.", "The precise problem")
text(s, 0.6, 1.8, 12.2, 4.8, [
    "• 47 reference monitors in Delhi (→ 60); planning norm = one per 25–50 km²",
    "• Site compliance = checklists: anti-smog guns per 5,000 m², CCTV, self-audit uploads",
    "• Flat fine per violation: ₹5 lakh for a large Delhi site, \"irrespective of the air quality category\" — doubled under GRAP",
    "• Oct 2025: tankers filmed spraying water at the Jahangirpuri reference monitor (CPCB says its data pipeline can't be manipulated)",
    ("Penalties track checklists, not the dust a site actually releases.", 22, ACCENT, True)], size=20, space=12)
source(s, "Sources: Hindustan Times (2026; Feb 2024); Newslaundry, 30 Oct 2025; New Indian Express, 3 Dec 2025. research/landscape.md §2")

# 4 — Who pays
s = slide("Who pays for not knowing: daily-wage workers", "The cost of not knowing lands on daily-wage workers. Delhi's welfare board has 1.38 million construction workers on its rolls, but only about 90 thousand active members. In one ground survey, 5 of 75 workers were registered and none had received the ₹8,000 aid. The Supreme Court had to order compensation upon GRAP-4.", "Socio-economic impact")
text(s, 0.6, 1.9, 6, 4, [("1,378,199", 44, ACCENT, True), ("construction workers on Delhi's welfare-board rolls — only 89,518 active members", 20, INK, False)])
text(s, 6.9, 1.9, 6, 4, [("5 of 75", 44, ACCENT, True), ("workers surveyed at sites were registered; none had received the ₹8,000 aid", 20, INK, False)])
text(s, 0.6, 4.9, 12, 1.2, [("Supreme Court (25 Nov 2024): assistance \"immediately … upon the imposition of GRAP-4\"", 18, GREY, False)])
source(s, "Source: Down To Earth (Dec 2024), citing the welfare-board dashboard and the SC order. research/landscape.md §2")

# 5 — Existing solutions
s = slide("Why today's tools fall short", "Indian compliance monitors stream a PM number to the DPCC cloud, but in Delhi's winter, background PM10 is around 300 micrograms. A sensor at the fence mostly measures Delhi, not the site. Anti-smog guns run on timers. DPCC told the NGT that smog towers are ineffective. IITM's forecasts are good at city scale but don't attribute sites. Global fence-line systems exist for mines and refineries.", "Background")
rows = [("Reference monitors (CAAQMS)", "city scale — can't attribute a site"),
        ("Compliance monitors → DPCC cloud", "absolute PM; site ≪ Delhi background in winter"),
        ("Anti-smog guns / sprinklers", "timer-driven; no feedback on whether dust fell"),
        ("Smog towers", "DPCC's own NGT affidavit: ineffective"),
        ("City AQI forecasts (IITM, Google)", "tomorrow's city AQI, not today's dusty site"),
        ("Fence-line systems (Aeroqual, Envirosuite)", "mines/refineries; not Indian construction workflows")]
tb = s.shapes.add_table(len(rows), 2, Inches(0.6), Inches(1.8), Inches(12.1), Inches(4.4)).table
for i, (a, b) in enumerate(rows):
    for j, v in enumerate((a, b)):
        c = tb.cell(i, j); c.text = v
        c.text_frame.paragraphs[0].font.size = Pt(18)
source(s, "Sources: research/landscape.md §3–4 (TERI winter PM10 314 ± 77 µg/m³; DPCC affidavit via Down To Earth; vendor pages)")

# 6 — Solution
s = slide("DustLedger: a fence-line ring that subtracts Delhi", "Four low-cost nodes on the boundary and a wind vane. Delhi's background reaches all four; the site's own plume reaches only the downwind ones. The difference is the site's contribution. When the wind is calm or confusing we say 'inconclusive' instead of guessing. When the site is the cause, misting turns on — not on a timer. Records are signed on the device and hash-chained, so edits become detectable. Fifteen-minute reporting with uncertainty is the design target; today's replay is hourly.", "Proposed solution")
s.shapes.add_picture(os.path.join(HERE, "fig_site.png"), Inches(0.5), Inches(1.6), height=Inches(5.2))
text(s, 6.6, 1.7, 6.3, 5.2, [
    ("1  Sense", 22, ACCENT, True), ("3–4 PM nodes + wind vane on the boundary; solar, LTE", 17, INK, False),
    ("2  Attribute", 22, ACCENT, True), ("downwind − upwind = the site's own dust; abstain in calm winds", 17, INK, False),
    ("3  Act", 22, ACCENT, True), ("mist only when the site is the cause (flow-metered)", 17, INK, False),
    ("4  Prove", 22, ACCENT, True), ("signed, hash-chained ledger → regulator ranking", 17, INK, False)], space=4)

# 7 — Evidence: forecasting
s = slide("Evidence on real Delhi data: forecasting isn't the gap", "We built the data pipeline on four Delhi monitors with two years of meteorology. Our first honest finding: with one year of history, machine learning did not beat yesterday's value for tomorrow's city PM. So we use forecasts only to schedule dusty work, and put the innovation in attribution. What's real today: the data pipeline and this benchmark. What's simulated: the site, its sensors and its plume.", "What we built")
s.shapes.add_picture(os.path.join(HERE, "fig_forecast.png"), Inches(0.5), Inches(1.6), height=Inches(5.2))
text(s, 10.6, 1.8, 2.4, 5, [("Real: CPCB/DPCC data (Anand Vihar, R K Puram, ITO, Major Dhyan Chand Stadium) + ERA5", 14, INK, False), ("Walk-forward, time-ordered test: Nov 2025–Sep 2026", 14, GREY, False)])

# 8 — Evidence: attribution replay
s = slide("Isolating one site's dust: the ring works where a fixed pair fails", "In a replay with real Delhi background and real wind, but a simulated site and sensors, a fixed upwind/downwind pair is badly wrong because the wind keeps changing direction. A wind-resolved ring cuts the error by about a quarter. Weekly reference visits barely help, because sensors foul and drift between visits. Re-calibrating the nodes against each other in the site's quiet hours cuts the error by another sixth, and that held up when an independent reviewer's harder stress tests were applied, including a site shape the method never saw. It's still a simulation, so we'll prove it on real hardware in November.", "The innovation")
s.shapes.add_picture(os.path.join(HERE, "fig_attribution.png"), Inches(0.5), Inches(1.6), height=Inches(5.2))
text(s, 10.6, 1.8, 2.4, 5, [("SIMULATED site, plume, sensors and faults on REAL background + wind", 14, WARN, True), ("Stress tests from an independent review: fouling, dust-vs-ambient gain, dirty nights, neighbours, vane bias", 14, INK, False)])

# 9 — Impact, buyers, scale
s = slide("Impact, buyers and scale — what we can honestly claim", "Today the buyer case is evidence and inspection triage. Exemptions for clean sites are a pilot we'll propose, not a promise. We won't quote an impact number until our November trial measures fence-line PM10 and water with misting on and off. The method subtracts the local background instead of modelling it, so it needs no city-specific forecast model; siting, local aerosol and calibration still need checking in each city.", "Impact & scalability")
text(s, 0.6, 1.8, 6.0, 5, [
    ("Who benefits", 22, ACCENT, True),
    "• Regulator: inspection triage for 334 enforcement teams",
    "• Builder: a measured record for disputes",
    "• Workers: proposed 'earned-permission' pilot for verified-clean sites",
    ("What we'll measure (Nov trial)", 22, ACCENT, True),
    "• Fence-line PM10 with misting on vs off",
    "• Litres of water per unit of dust avoided"], size=18)
text(s, 6.9, 1.8, 6.0, 5, [
    ("Scale across India", 22, ACCENT, True),
    "• NCAP: only Delhi, Pune, Mumbai and Ahmedabad have a decision-support system",
    "• Differential method: no city-specific forecast model",
    "• Same hardware; local siting + calibration checks",
    "• Plugs into existing state portals (DPCC cloud pattern)"], size=18)
source(s, "Sources: The Hindu (2024) on NCAP; The News Mill (Sept 2026) on enforcement teams. research/landscape.md")

# 10 — Roadmap + ask
s = slide("Roadmap to a working prototype — and our ask", "By December: a working four-node ring with closed-loop misting, tested on a controlled dust source and, with a partner, at a real site — with measured numbers. Measure the site, not the city, and let clean sites work.", "Roadmap")
text(s, 0.6, 1.8, 12.2, 4.5, [
    "Oct  — evaluator v2 stress tests; spike ratchet; parts ordered",
    "Nov (1st half) — 4-node ring + wind + misting rig; side-by-side and reference co-location; controlled dust trial",
    "1 Dec — report + working prototype (real sensor data, labelled)",
    "15–18 Dec — finals demo: live ring around a dust source, dashboard, automatic misting",
    ("Ask: a pilot construction site · a DPCC/CAQM contact · a reference co-location slot", 22, ACCENT, True)], size=20, space=12)
text(s, 0.6, 6.2, 12, 0.6, [("Measure the site, not the city — and let clean sites work.", 24, INK, True)])

# Appendix A — time series
s = slide("Appendix A — one week of the replay", "Backup slide: a week in November 2025. Top: the real Delhi background PM10 at a CPCB/DPCC monitor. Bottom: a simulated site's true contribution versus the fixed-pair estimate and our estimate. Gaps are 'inconclusive' calm-wind hours.", "Backup")
s.shapes.add_picture(os.path.join(HERE, "fig_timeseries.png"), Inches(0.6), Inches(1.5), height=Inches(5.4))

# Appendix B — what is real vs simulated
s = slide("Appendix B — what is real, what is simulated", "Backup slide for the 'what have you built?' question.", "Backup")
text(s, 0.6, 1.8, 6, 5, [("Real (committed, reproducible)", 22, ACCENT, True), "• CPCB/DPCC hourly PM at 4 Delhi monitors (OpenAQ archive)", "• ERA5 meteorology; CAMS fields (reference)", "• Forecast benchmark, walk-forward split", "• Typed component contracts, replay harness"], size=18)
text(s, 6.9, 1.8, 6, 5, [("Simulated (labelled)", 22, WARN, True), "• Construction site, emissions, plumes", "• Fence-line sensors and their errors", "• Misting effect: not modelled yet", ("Not yet built", 22, GREY, True), "• Hardware ring, misting rig, tamper detection, dashboard"], size=18)

prs.save(os.path.join(HERE, "DustLedger_zonal.pptx"))
print("deck written:", len(prs.slides), "slides")
